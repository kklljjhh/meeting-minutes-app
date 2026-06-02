from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from datetime import datetime

def generate_meeting_minutes(transcription_text, meeting_topic="会议纪要"):
    doc = Document()
    
    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font.size = Pt(12)
    
    title = doc.add_heading(level=1)
    title.text = meeting_topic
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    date_str = datetime.now().strftime('%Y年%m月%d日')
    date_paragraph = doc.add_paragraph(f"日期：{date_str}")
    date_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    doc.add_heading('一、会议内容', level=2)
    
    lines = transcription_text.split('\n')
    for line in lines:
        line = line.strip()
        if line:
            doc.add_paragraph(line, style='List Paragraph')
    
    doc.add_heading('二、会议决议', level=2)
    doc.add_paragraph('1. 根据会议讨论内容，形成以下决议：')
    doc.add_paragraph('2. 相关责任人需按照决议内容执行。')
    
    doc.add_heading('三、下次会议安排', level=2)
    doc.add_paragraph('待定，将另行通知。')
    
    return doc

def save_document(doc, output_path):
    doc.save(output_path)

if __name__ == '__main__':
    sample_text = """参会人员讨论了项目进度问题。
张三汇报了当前开发进度。
李四提出了改进建议。
王总做了总结发言。"""
    
    doc = generate_meeting_minutes(sample_text, "LDAR社区平台升级与运营策略会议纪要")
    save_document(doc, "test_output.docx")
    print("文档已生成")
